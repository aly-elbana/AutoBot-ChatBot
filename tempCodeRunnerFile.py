import gradio as gr
import os
from dotenv import load_dotenv
from PyPDF2 import PdfReader
from pptx import Presentation
from langchain_google_genai import ChatGoogleGenerativeAI

load_dotenv()
api_key = os.getenv("GOOGLE_API_KEY")

llm = ChatGoogleGenerativeAI(model="gemini-1.5-flash", google_api_key=api_key)


def read_pdf(file_path):
    text = ""
    reader = PdfReader(file_path)
    for page in reader.pages:
        text += page.extract_text() + "\n"
    return text

def read_pptx(file_path):
    text = ""
    prs = Presentation(file_path)
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text

def read_text(file_path):
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def process_file(file):
    if not file:
        return None, None
    
    file_path = file.name
    ext = os.path.splitext(file_path)[-1].lower()

    if ext == ".pdf":
        return read_pdf(file_path), "text"
    elif ext == ".pptx":
        return read_pptx(file_path), "text"
    elif ext in [".txt", ".py", ".md", ".json"]:
        return read_text(file_path), "text"
    elif ext in [".png", ".jpg", ".jpeg", ".webp"]:
        return file_path, "image"
    else:
        return "❌ Unsupported file format", "error"


def chat(message, history, file):
    content, file_type = process_file(file) if file else (None, None)

    if file_type == "text" and content:
        user_input = f"{message}\n\n--- Attached File Content ---\n{content}"
        response = llm.invoke(user_input)
        return response.content

    elif file_type == "image" and content:
        response = llm.invoke([{"role": "user", "content": [{"type": "text", "text": message},
                                                            {"type": "image_url", "image_url": f"data:image/jpeg;base64,{content}"}]}])
        return response.content

    elif file_type == "error":
        return content

    else:
        response = llm.invoke(message)
        return response.content


with gr.Blocks(css="""
    .file-upload-box {
        height: 60px !important;
        max-height: 60px !important;
    }
""") as demo:
    chatbot = gr.Chatbot(height=400)
    with gr.Row():
        msg = gr.Textbox(placeholder="Type your question here...", scale=8)
        file_upload = gr.File(label="Upload", file_types=[".pdf", ".pptx", ".txt", ".py", ".md", ".json", ".png", ".jpg", ".jpeg", ".webp"], scale=2, elem_classes="file-upload-box")
        send = gr.Button("Send", scale=1)
    with gr.Row():
        clear = gr.Button("Clear Chat", scale=2)

    def respond(message, history, file):
        bot_msg = chat(message, history, file)
        history.append((message, bot_msg))
        return "", history

    send.click(respond, [msg, chatbot, file_upload], [msg, chatbot])
    msg.submit(respond, [msg, chatbot, file_upload], [msg, chatbot])
    clear.click(lambda: None, None, chatbot, queue=False)

demo.launch()
